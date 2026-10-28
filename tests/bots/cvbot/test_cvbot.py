from __future__ import annotations

import asyncio
import io
import json
from datetime import UTC, datetime, timedelta

import pytest
from aiogram.types import Chat, Document, InlineQuery, Message, RefundedPayment, Update
from sqlalchemy import event, func, select
from sqlalchemy.dialects import postgresql
from sqlalchemy.ext.asyncio import AsyncSession

from bots.cvbot.extraction import ExtractionError, extract, tesseract_text
from bots.cvbot.models import CVInvoice, CVReviewCredit
from bots.cvbot.payments import ProPayload
from bots.cvbot.pdf import render_pdf
from bots.cvbot.schemas import Analysis, ImprovedCV
from core.ai.base import AIError
from core.db.models import Payment
from core.runtime.dispatcher import build_dispatcher
from core.sdk.registry import discover
from tests.conftest import build_test_ctx, tg_user
from tests.fake_ai import fake_ai

SOURCE = "Alice Example\nalice@example.com\nSoftware engineer\nBuilt APIs in Python and SQL.\nEducation: BSc."
REPORT = {
    "score": 82, "subscores": {"ats": 80, "impact": 75, "clarity": 90, "formatting": 83},
    "strengths": ["Clear technical skills"], "issues": [{"issue": "No measured outcomes", "fix": "Add verified metrics"}],
    "rewritten_bullets": [{"original": "Built <APIs>", "improved": "Built reliable APIs using Python"}],
    "missing_sections": ["Projects"], "suggested_roles": ["Backend engineer"], "language": "English",
}
IMPROVED = {
    "name": "Alice Example", "headline": "Software engineer", "contact": "alice@example.com",
    "summary": "Engineer building reliable APIs.",
    "sections": [{"title": "Experience", "bullets": ["Built reliable APIs using Python."]}],
}


@pytest.fixture
async def cctx(sm, redis, api, tmp_path):
    ai, provider = fake_ai(redis, REPORT)
    ctx = await build_test_ctx(discover(["cvbot"])[0], sm, redis, api, tmp_path, ai=ai)
    ctx.extras["fake_provider"] = provider
    return ctx


async def unthrottle(ctx):
    keys = await ctx.redis.keys(ctx.keys("throttle*"))
    if keys:
        await ctx.redis.delete(*keys)


def fixture_pdf(text=SOURCE, pages=1):
    import pymupdf as fitz

    with fitz.open() as document:
        for _ in range(pages):
            page = document.new_page()
            page.insert_text((50, 70), text)
        return document.tobytes()


def test_extract_real_pdf_docx_image_and_caps():
    from docx import Document as WordDocument
    from PIL import Image

    assert "Software engineer" in extract(fixture_pdf(), "cv.pdf").text
    doc = WordDocument()
    doc.add_paragraph(SOURCE)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Python"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract(buf.getvalue(), "cv.docx").text
    assert "alice@example.com" in text and "Python" in text
    image_buf = io.BytesIO()
    Image.new("RGB", (300, 200), "white").save(image_buf, "PNG")
    assert len(extract(image_buf.getvalue(), "cv.png").images) == 1
    scanned = extract(fixture_pdf(""), "scan.pdf")
    assert not scanned.text and len(scanned.images) == 1
    with pytest.raises(ExtractionError):
        extract(fixture_pdf(pages=3), "cv.pdf", max_pages=2)
    with pytest.raises(ExtractionError):
        extract(fixture_pdf(), "cv.pdf", max_chars=10)
    with pytest.raises(ExtractionError):
        extract(b"not a PDF", "cv.pdf")


@pytest.mark.parametrize("template", ["classic", "modern"])
def test_pdf_templates_include_verified_content(template):
    import pymupdf as fitz

    data = render_pdf(ImprovedCV.model_validate(IMPROVED), template)
    with fitz.open(stream=data, filetype="pdf") as document:
        assert len(document) == 1
        text = "".join(page.get_text() for page in document)
        assert "Alice Example" in text and "Built reliable APIs" in text


async def test_real_document_update_renders_report_and_stores_only_summary(cctx, api, monkeypatch):
    async def download(*args, **kwargs):
        return fixture_pdf()

    monkeypatch.setattr("bots.cvbot.handlers.fetch_bytes", download)
    dp = build_dispatcher(cctx)
    update = Update(update_id=900, message=Message(
        message_id=900, date=datetime.now(UTC), chat=Chat(id=42, type="private"), from_user=tg_user(),
        document=Document(file_id="cv_pdf", file_unique_id="unique", file_name="cv.pdf", file_size=1000),
    ))
    await dp.feed_update(cctx.bot, update)
    texts = [m.text for m in api.calls("SendMessage")]
    assert any("82/100" in t for t in texts)
    assert any("Built &lt;APIs&gt;" in t for t in texts)
    report = await cctx.extras["cv"].latest(42)
    assert report.score == 82 and "rewritten_bullets" not in report.summary
    assert "alice@example.com" not in json.dumps(report.summary)
    assert await cctx.extras["cv_quota"].used(42) == 1
    assert (await cctx.users.get(42)).actions_count == 1
    request = cctx.extras["fake_provider"].requests[-1]
    assert request.schema and request.schema["additionalProperties"] is False


async def test_quota_vip_and_failed_analysis_refund(cctx, api, builders):
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    await unthrottle(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    assert api.calls("SendMessage")[-1].text == cctx.translator.get("en", "cv-quota", max=1)
    assert len(cctx.extras["fake_provider"].requests) == 1
    await cctx.subs.grant_vip(42, 2)
    await unthrottle(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    assert len(cctx.extras["fake_provider"].requests) == 2
    assert await cctx.extras["cv_quota"].used(42) == 1

    cctx.extras["fake_provider"].reply = AIError("fake failure")
    await dp.feed_update(cctx.bot, builders["message"](SOURCE, user_id=43))
    assert await cctx.extras["cv_quota"].used(43) == 0
    assert await cctx.extras["cv"].latest(43) is None


async def test_scan_quota_reserved_before_vision(cctx, api, builders, monkeypatch):
    async def download(*args, **kwargs):
        return fixture_pdf("")

    monkeypatch.setattr("bots.cvbot.handlers.fetch_bytes", download)
    await cctx.users.get_or_create(42)
    await cctx.extras["cv_quota"].consume(42, 1)
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, Update(update_id=901, message=Message(
        message_id=901, date=datetime.now(UTC), chat=Chat(id=42, type="private"), from_user=tg_user(),
        document=Document(file_id="scan", file_unique_id="scan", file_name="scan.pdf", file_size=1000),
    )))
    assert cctx.extras["fake_provider"].requests == []
    assert "free full analysis" in api.calls("SendMessage")[-1].text


async def test_forget_scopes_owner_and_clears_session_and_fsm(cctx, api, builders):
    dp = build_dispatcher(cctx)
    for user_id in (42, 43):
        await dp.feed_update(cctx.bot, builders["message"](SOURCE, user_id=user_id))
    await unthrottle(cctx)
    await dp.feed_update(cctx.bot, builders["message"]("/match"))
    await unthrottle(cctx)
    await dp.feed_update(cctx.bot, builders["message"]("/forget"))
    assert await cctx.extras["cv"].latest(42) is None
    assert cctx.extras["cv"].session(42) is None
    assert await cctx.extras["cv"].latest(43) is not None
    for key in await cctx.redis.keys("*"):
        if await cctx.redis.type(key) == b"string":
            value = await cctx.redis.get(key)
            assert b"alice@example.com" not in value
    # A request started before /forget is forbidden from restoring its report.
    assert await cctx.extras["cv"].save(42, Analysis.model_validate(REPORT), epoch=0) is None


async def test_score_card_and_inline_are_personal(cctx, api, builders):
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    await unthrottle(cctx)
    await dp.feed_update(cctx.bot, builders["message"]("/share"))
    photo = api.calls("SendPhoto")[-1]
    assert photo.photo.data.startswith(b"\x89PNG") and "https://t.me/test_bot" in photo.caption
    await dp.feed_update(cctx.bot, Update(update_id=902, inline_query=InlineQuery(
        id="inline", from_user=tg_user(), query="", offset="",
    )))
    inline = api.calls("AnswerInlineQuery")[-1]
    assert inline.is_personal and inline.cache_time == 0 and len(inline.results) == 1
    await dp.feed_update(cctx.bot, Update(update_id=903, inline_query=InlineQuery(
        id="inline_other", from_user=tg_user(44), query="", offset="",
    )))
    assert api.calls("AnswerInlineQuery")[-1].results == []


async def test_mock_interview_five_turns_and_arabic(cctx, api, builders):
    def reply(request):
        title = request.schema["title"]
        if title == "Analysis":
            return REPORT
        if title == "InterviewQuestions":
            return {"questions": [f"Question {i}?" for i in range(1, 6)]}
        if title == "AnswerFeedback":
            return {"score": 70, "feedback": "Explain the outcome clearly."}
        raise AssertionError(title)

    cctx.extras["fake_provider"].reply = reply
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE, lang="ar"))
    for text in ["/interview", "Backend engineer", *["I built a Python service." for _ in range(5)]]:
        await unthrottle(cctx)
        await dp.feed_update(cctx.bot, builders["message"](text, lang="ar"))
    assert api.calls("SendMessage")[-1].text == cctx.translator.get("ar", "cv-interview-done")
    requests = cctx.extras["fake_provider"].requests
    assert len(requests) == 7 and all(r.fast for r in requests[1:])
    assert "in ar" in requests[-1].system
    assert cctx.extras["cv"].session(42).questions == []


async def test_pdf_export_enforces_vip_and_uses_selected_template(cctx, api, builders):
    cctx.extras["fake_provider"].reply = lambda r: IMPROVED if r.schema["title"] == "ImprovedCV" else REPORT
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    await dp.feed_update(cctx.bot, builders["callback"]("cv:pdf:classic"))
    assert not api.calls("SendDocument")
    await cctx.subs.grant_vip(42, 2)
    await unthrottle(cctx)
    await dp.feed_update(cctx.bot, builders["callback"]("cv:pdf:modern"))
    assert api.calls("SendDocument")[-1].document.filename == "cv-modern.pdf"


async def test_pro_invoice_routing_validation_and_exactly_once_credit(cctx, api, builders):
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, builders["message"]("/pro"))
    invoice = api.calls("SendInvoice")[-1]
    payload = ProPayload.decode(invoice.payload)
    assert payload and payload.stars == 25
    for raw, amount, user_id, currency in [
        (invoice.payload, 25, 43, "XTR"), (invoice.payload, 26, 42, "XTR"),
        (invoice.payload, 25, 42, "USD"), (invoice.payload.replace("cvbot", "notes"), 25, 42, "XTR"),
        (ProPayload("cvbot", 42, 25, "x" * 22).encode(), 25, 42, "XTR"),
    ]:
        await dp.feed_update(cctx.bot, builders["pre_checkout"](raw, amount, user_id=user_id, currency=currency))
        assert api.calls("AnswerPreCheckoutQuery")[-1].ok is False
    await dp.feed_update(cctx.bot, builders["pre_checkout"](invoice.payload, 25))
    assert api.calls("AnswerPreCheckoutQuery")[-1].ok is True
    await dp.feed_update(cctx.bot, builders["payment"](invoice.payload, 25, "wrong-payer", user_id=43))
    await dp.feed_update(cctx.bot, builders["payment"](invoice.payload, 26, "wrong-amount"))
    await dp.feed_update(cctx.bot, builders["payment"](invoice.payload, 25, "pro-charge"))
    await dp.feed_update(cctx.bot, builders["payment"](invoice.payload, 25, "pro-charge"))
    await dp.feed_update(cctx.bot, builders["payment"](invoice.payload, 25, "same-invoice-other-charge"))
    async with cctx.sm() as session:
        assert (await session.execute(select(func.count(CVReviewCredit.payment_id)))).scalar_one() == 1
        assert (await session.execute(select(func.count(Payment.id)))).scalar_one() == 1
    assert not (await cctx.users.get(42)).is_vip()
    await unthrottle(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    assert cctx.extras["fake_provider"].requests[-1].max_tokens == 6000
    assert await cctx.extras["cv_quota"].used(42) == 0
    assert await cctx.extras["cv_pro"].consume(42) is None


async def test_pro_expiry_failure_restore_and_refund_preserves_vip(cctx, api, builders):
    dp = build_dispatcher(cctx)
    await cctx.users.get_or_create(42)
    reviews = cctx.extras["cv_pro"]
    expired = await reviews.issue(42, 25)
    async with cctx.sm() as session:
        row = await session.get(CVInvoice, expired.nonce)
        row.expires_at = datetime.now(UTC) - timedelta(seconds=1)
        await session.commit()
    await dp.feed_update(cctx.bot, builders["pre_checkout"](expired.encode(), 25))
    assert api.calls("AnswerPreCheckoutQuery")[-1].ok is False
    payload = await reviews.issue(42, 25)
    await dp.feed_update(cctx.bot, builders["payment"](payload.encode(), 25, "refund-charge"))
    cctx.extras["fake_provider"].reply = AIError("unavailable")
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    credit = await reviews.consume(42)
    assert credit is not None
    await reviews.restore(42, credit)
    until = await cctx.subs.grant_vip(42, 10)
    refund = RefundedPayment(currency="XTR", total_amount=25, invoice_payload=payload.encode(),
                             telegram_payment_charge_id="refund-charge")
    await dp.feed_update(cctx.bot, Update(update_id=904, message=Message(
        message_id=904, date=datetime.now(UTC), chat=Chat(id=42, type="private"), from_user=tg_user(),
        refunded_payment=refund,
    )))
    assert await reviews.consume(42) is None
    assert (await cctx.users.get(42)).vip_until == until


async def test_unconfigured_ai_never_downloads_or_consumes(cctx, api, builders):
    cctx.ai = None
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, builders["message"](SOURCE))
    assert api.calls("SendMessage")[-1].text == cctx.translator.get("en", "ai-unavailable")
    assert await cctx.extras["cv_quota"].used(42) == 0
    assert await cctx.extras["cv"].latest(42) is None


async def test_job_match_cover_letter_and_linkedin_flows(cctx, api, builders):
    def reply(request):
        if request.schema is None:
            return "Dear Hiring Manager,\nI build Python APIs. <Thank you>"
        return {
            "Analysis": REPORT,
            "JobMatch": {"score": 71, "missing_keywords": ["Kubernetes"], "tailored_bullets": ["Built Python APIs"]},
            "LinkedIn": {"headline": "Python <Engineer>", "summary": "I build reliable APIs."},
        }[request.schema["title"]]

    cctx.extras["fake_provider"].reply = reply
    dp = build_dispatcher(cctx)
    for text in [SOURCE, "/match", "Backend developer, Python, Kubernetes", "/coverletter", "-", "/linkedin"]:
        await unthrottle(cctx)
        await dp.feed_update(cctx.bot, builders["message"](text))
    messages = [m.text for m in api.calls("SendMessage")]
    assert any("71/100" in text and "Kubernetes" in text for text in messages)
    assert any("&lt;Thank you&gt;" in text for text in messages)
    assert "Python &lt;Engineer&gt;" in messages[-1]
    assert await cctx.extras["cv_quota"].used(42) == 1


async def test_scanned_cv_vision_then_analysis_in_memory(cctx, api, monkeypatch):
    async def download(*args, **kwargs):
        return fixture_pdf("")

    monkeypatch.setattr("bots.cvbot.handlers.fetch_bytes", download)
    cctx.extras["fake_provider"].reply = lambda r: SOURCE if r.attachments else REPORT
    dp = build_dispatcher(cctx)
    await dp.feed_update(cctx.bot, Update(update_id=905, message=Message(
        message_id=905, date=datetime.now(UTC), chat=Chat(id=42, type="private"), from_user=tg_user(),
        document=Document(file_id="scan", file_unique_id="scan", file_name="scan.pdf", file_size=1000),
    )))
    requests = cctx.extras["fake_provider"].requests
    assert len(requests) == 2 and requests[0].attachments[0].kind == "image"
    assert requests[1].attachments == ()
    session = cctx.extras["cv"].session(42)
    assert session.source == SOURCE and isinstance(session.source, str)


def test_offline_ocr_deletes_temporary_image(monkeypatch):
    import subprocess
    from pathlib import Path

    seen = []
    monkeypatch.setattr("bots.cvbot.extraction.shutil.which", lambda name: "tesseract")

    def run(args, **kwargs):
        seen.append(Path(args[1]))
        assert seen[-1].read_bytes() == b"image"
        assert "ara+eng" in args and kwargs["timeout"] == 30
        return subprocess.CompletedProcess(args, 0, stdout=SOURCE.encode(), stderr=b"")

    monkeypatch.setattr("bots.cvbot.extraction.subprocess.run", run)
    assert tesseract_text(b"image") == SOURCE
    assert not seen[0].exists() and not seen[0].parent.exists()


async def test_sources_actually_expire_without_another_update(cctx, monkeypatch):
    monkeypatch.setattr("bots.cvbot.service.SESSION_SECONDS", 0.01)
    service = cctx.extras["cv"]
    service.remember(42, SOURCE, Analysis.model_validate(REPORT))
    assert service._sessions[42].source == SOURCE
    await asyncio.sleep(0.05)
    # No session()/prune()/scheduler call: the gateway timer removes sensitive source memory itself.
    assert service._sessions == {} and service._expiry_handles == {}


async def test_replaced_evicted_and_forgotten_sources_cancel_timers(cctx, monkeypatch):
    monkeypatch.setattr("bots.cvbot.service.MAX_SESSIONS", 2)
    service = cctx.extras["cv"]
    report = Analysis.model_validate(REPORT)
    service.remember(42, SOURCE, report)
    original = service._sessions[42]
    old_handle = service._expiry_handles[42]
    service.remember(42, SOURCE + "Updated", report)
    assert old_handle.cancelled()
    service._expire(42, original.generation)
    assert service._sessions[42].source.endswith("Updated")
    service.remember(43, SOURCE, report)
    service.remember(44, SOURCE, report)
    assert set(service._sessions) == set(service._expiry_handles) == {43, 44}
    handle = service._expiry_handles[43]
    await service.forget(43)
    assert handle.cancelled() and 43 not in service._expiry_handles


async def test_privacy_locks_only_owner_table_on_postgresql(cctx):
    """BotUser eager-loads profile with an OUTER JOIN; PostgreSQL rejects an unrestricted row lock."""
    statements = []

    def record(orm_state):
        if orm_state.is_select and getattr(orm_state.statement, "_for_update_arg", None) is not None:
            statements.append(str(orm_state.statement.compile(dialect=postgresql.dialect())))

    await cctx.users.get_or_create(42)
    event.listen(AsyncSession.sync_session_class, "do_orm_execute", record)
    try:
        await cctx.extras["cv"].save(42, Analysis.model_validate(REPORT), epoch=0)
        await cctx.extras["cv"].forget(42)
    finally:
        event.remove(AsyncSession.sync_session_class, "do_orm_execute", record)
    assert len(statements) == 2
    assert all("LEFT OUTER JOIN users" in stmt and stmt.endswith("FOR UPDATE OF bot_users") for stmt in statements)
