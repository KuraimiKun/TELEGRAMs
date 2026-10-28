import zipfile

import pymupdf as fitz
import pytest
from docx import Document
from PIL import Image, ImageDraw

from bots.doctools.operations import DocumentError, extract_text, parse_ranges, process, tesseract_text


def pdf_file(path, texts):
    with fitz.open() as doc:
        for text in texts:
            page = doc.new_page()
            page.insert_text((30, 40), text)
        doc.save(path)
    return path


def test_merge_preserves_file_and_page_order(tmp_path):
    first = pdf_file(tmp_path / "first.pdf", ["first page", "second page"])
    second = pdf_file(tmp_path / "second.pdf", ["third page"])
    result = process("merge", [first, second], tmp_path)
    assert result.pages == 3
    with fitz.open(result.files[0]) as doc:
        assert [page.get_text().strip() for page in doc] == ["first page", "second page", "third page"]


def test_split_extract_and_compress_preserve_contents(tmp_path):
    source = pdf_file(tmp_path / "source.pdf", ["one", "two", "three", "four"])
    split = process("split", [source], tmp_path, options={"ranges": "1-2;3-4"})
    assert len(split.files) == 2
    with fitz.open(split.files[1]) as doc:
        assert [p.get_text().strip() for p in doc] == ["three", "four"]
    extract = process("extract", [source], tmp_path, options={"ranges": "4,2"})
    with fitz.open(extract.files[0]) as doc:
        assert [p.get_text().strip() for p in doc] == ["four", "two"]
    compressed = process("compress", [source], tmp_path)
    with fitz.open(compressed.files[0]) as doc:
        assert len(doc) == 4 and doc[0].get_text().strip() == "one"


@pytest.mark.parametrize("ranges", ["", "0", "5", "3-1", "1,1", "1-2,2", "x", "1;2"])
def test_page_range_validation(ranges):
    with pytest.raises(DocumentError, match="pages"):
        parse_ranges(ranges, 4)


def test_page_caps_enforced_before_merge(tmp_path):
    paths = [pdf_file(tmp_path / f"{i}.pdf", ["one", "two"]) for i in range(2)]
    with pytest.raises(DocumentError, match="pages"):
        process("merge", paths, tmp_path, max_pages=3)
    assert not (tmp_path / "merged.pdf").exists()


def test_split_output_page_cap_prevents_range_expansion(tmp_path):
    source = pdf_file(tmp_path / "source.pdf", ["one", "two", "three"])
    with pytest.raises(DocumentError, match="pages"):
        process("split", [source], tmp_path, options={"ranges": "1-3;1-3"}, max_pages=3)
    assert not (tmp_path / "split-001.pdf").exists()


def test_images_to_pdf_are_ordered_and_a4(tmp_path):
    paths = []
    for i, color in enumerate(["red", "blue"]):
        path = tmp_path / f"image{i}.png"
        Image.new("RGB", (300, 100), color).save(path)
        paths.append(path)
    result = process("image_pdf", paths, tmp_path)
    with fitz.open(result.files[0]) as doc:
        assert len(doc) == 2
        assert doc[0].rect.width == pytest.approx(595.28, abs=0.01)
        assert doc[0].rect.height == pytest.approx(841.89, abs=0.01)
        for page, expected in zip(doc, [0, 2], strict=True):
            pixmap = page.get_pixmap()
            center = pixmap.pixel(pixmap.width // 2, pixmap.height // 2)
            assert center[expected] > 200


def test_rotate_numbers_watermark_and_password_roundtrip(tmp_path):
    source = pdf_file(tmp_path / "source.pdf", ["Document content"])
    rotated = process("rotate", [source], tmp_path, options={"angle": 90})
    with fitz.open(rotated.files[0]) as doc:
        assert doc[0].rotation == 90
    numbered = process("numbers", [source], tmp_path)
    with fitz.open(numbered.files[0]) as doc:
        assert "1" in doc[0].get_text()
    watermark = process("watermark", [source], tmp_path, options={"watermark": "Study copy"})
    with fitz.open(watermark.files[0]) as doc:
        assert "Study copy" in doc[0].get_text()
    locked = process("encrypt", [source], tmp_path, options={"password": "secret"})
    with fitz.open(locked.files[0]) as doc:
        assert doc.needs_pass
    with pytest.raises(DocumentError, match="password"):
        process("decrypt", locked.files, tmp_path, options={"password": "wrong"})
    unlocked = process("decrypt", locked.files, tmp_path, options={"password": "secret"})
    with fitz.open(unlocked.files[0]) as doc:
        assert not doc.needs_pass and "Document content" in doc[0].get_text()


def test_pdf_images_zip_and_scan_enhancement(tmp_path):
    source = pdf_file(tmp_path / "source.pdf", ["one", "two"])
    result = process("images", [source], tmp_path)
    with zipfile.ZipFile(result.files[0]) as archive:
        assert archive.namelist() == ["page-0001.png", "page-0002.png"]
    image = Image.new("RGB", (300, 150), "#eeeeee")
    ImageDraw.Draw(image).text((20, 50), "A short study note", fill="#777777")
    image = image.rotate(3, fillcolor="white")
    path = tmp_path / "scan.png"
    image.save(path)
    enhanced = process("enhance", [path], tmp_path)
    with Image.open(enhanced.files[0]) as output:
        assert output.mode == "L"
        assert output.getextrema()[0] < 100


def test_counts_include_docx_tables_and_arabic_text(tmp_path):
    source = tmp_path / "assignment.docx"
    document = Document()
    document.add_paragraph("Hello world")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "مرحبا بالعالم"
    document.save(source)
    counts = extract_text([source], max_pages=10)
    assert counts.words == 4
    assert counts.chars == len("Hello world\nمرحبا بالعالم")


def test_compress_downscales_embedded_images(tmp_path):
    import io

    image = Image.new("RGB", (3000, 2000), "#aaaaaa")
    ImageDraw.Draw(image).rectangle((10, 10, 2000, 1000), fill="black")
    stream = io.BytesIO()
    image.save(stream, "PNG")
    source = tmp_path / "photo.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_image(page.rect, stream=stream.getvalue())
        doc.save(source)
    output = process("compress", [source], tmp_path).files[0]
    assert output.stat().st_size < source.stat().st_size
    with fitz.open(output) as doc:
        info = doc[0].get_images()[0]
        assert max(info[2], info[3]) <= 1600


def test_tesseract_fallback_requests_arabic_and_english(tmp_path, monkeypatch):
    from types import SimpleNamespace

    calls = []
    monkeypatch.setattr("bots.doctools.operations.shutil.which", lambda command: "tesseract")

    def run(args, **kwargs):
        calls.append((args, kwargs))
        return SimpleNamespace(stdout="مرحبا English".encode())

    monkeypatch.setattr("bots.doctools.operations.subprocess.run", run)
    assert tesseract_text([tmp_path / "scan.png"], timeout=10) == "مرحبا English"
    assert calls[0][0][-4:] == ["-l", "ara+eng", "--psm", "3"]
    assert calls[0][1]["timeout"] == 10


def test_missing_ocr_backend_has_user_facing_error(tmp_path, monkeypatch):
    monkeypatch.setattr("bots.doctools.operations.shutil.which", lambda command: None)
    with pytest.raises(DocumentError, match="unavailable"):
        tesseract_text([tmp_path / "scan.png"])
