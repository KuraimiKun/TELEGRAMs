"""Bounded document transformations, invoked in worker threads only.

All outputs live in the caller's private scratch directory. Source paths are never
derived from user filenames, and no PDF JavaScript or embedded files are executed.
"""

from __future__ import annotations

import io
import itertools
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pymupdf as fitz
from docx import Document
from PIL import Image, ImageEnhance, ImageOps, UnidentifiedImageError

MAX_IMAGE_PIXELS = 25_000_000
MAX_RENDER_PIXELS = 12_000_000
MAX_TEXT_CHARS = 200_000


class DocumentError(ValueError):
    def __init__(self, code: str = "invalid") -> None:
        self.code = code
        super().__init__(code)


@dataclass(slots=True)
class OperationResult:
    files: list[Path] = field(default_factory=list)
    text: str = ""
    pages: int = 0
    words: int = 0
    chars: int = 0


def parse_ranges(value: str, page_count: int) -> list[int]:
    """Parse 1-based selections, preserving order and rejecting duplicates or bad bounds."""
    if not value.strip() or len(value) > 2000:
        raise DocumentError("pages")
    pages: list[int] = []
    for part in value.split(","):
        match = re.fullmatch(r"\s*(\d+)(?:\s*-\s*(\d+))?\s*", part)
        if not match:
            raise DocumentError("pages")
        first = int(match[1])
        last = int(match[2] or first)
        if not 1 <= first <= last <= page_count:
            raise DocumentError("pages")
        pages.extend(range(first - 1, last))
        if len(pages) > page_count:
            raise DocumentError("pages")
    if len(pages) != len(set(pages)):
        raise DocumentError("pages")
    return pages


def _pdf(path: Path, password: str = "") -> Any:
    try:
        doc = fitz.open(path)
    except (RuntimeError, ValueError) as e:
        raise DocumentError() from e
    if not doc.is_pdf:
        doc.close()
        raise DocumentError()
    if doc.needs_pass and not doc.authenticate(password):
        doc.close()
        raise DocumentError("password")
    return doc


def _image(path: Path) -> Image.Image:
    try:
        with Image.open(path) as source:
            if source.width * source.height > MAX_IMAGE_PIXELS:
                raise DocumentError("large")
            image = ImageOps.exif_transpose(source).convert("RGB")
            image.load()
            return image
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as e:
        raise DocumentError() from e


def _render(page: Any, scale: float = 1.5) -> bytes:
    rect = page.rect
    if rect.width <= 0 or rect.height <= 0:
        raise DocumentError()
    # Adversarial PDFs can declare enormous pages; cap pixels before allocating.
    scale = min(scale, (MAX_RENDER_PIXELS / (rect.width * rect.height)) ** 0.5)
    return page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes("png")


def images_to_pdf(paths: list[Path], output: Path) -> Path:
    """One image per A4 page, in arrival order; fit inside a 24pt margin."""
    with fitz.open() as result:
        for path in paths:
            image = _image(path)
            image.thumbnail((2400, 3400))
            stream = io.BytesIO()
            image.save(stream, format="JPEG", quality=90)
            page = result.new_page(width=595.28, height=841.89)
            page.insert_image(fitz.Rect(24, 24, 571.28, 817.89), stream=stream.getvalue(), keep_proportion=True)
        result.save(output, garbage=4, deflate=True)
    return output


def _projection_score(image: Image.Image, angle: float) -> float:
    rotated = image.rotate(angle, expand=False, fillcolor=255)
    tiny = rotated.resize((160, rotated.height))
    pixels = tiny.tobytes()
    sums = [sum(255 - p for p in pixels[y * 160:(y + 1) * 160]) for y in range(tiny.height)]
    return float(sum((b - a) ** 2 for a, b in itertools.pairwise(sums)))


def enhance_scan(path: Path, output: Path) -> Path:
    image = ImageOps.grayscale(_image(path))
    image = ImageOps.autocontrast(image, cutoff=1)
    sample = image.copy()
    sample.thumbnail((400, 500))
    # Search a bounded +/-5 degree range; projection peaks at aligned text rows.
    angles = [i / 2 for i in range(-10, 11)]
    best = max(angles, key=lambda angle: _projection_score(sample, angle))
    baseline = _projection_score(sample, 0)
    if best and _projection_score(sample, best) > baseline * 1.05:
        image = image.rotate(best, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=255)
    image = ImageEnhance.Contrast(image).enhance(1.35)
    image.save(output, "PNG", optimize=True)
    return output


def text_counts(text: str, pages: int = 0) -> OperationResult:
    return OperationResult(text=text, pages=pages, words=len(re.findall(r"\b\w+\b", text)), chars=len(text))


def extract_text(paths: list[Path], *, max_pages: int, password: str = "") -> OperationResult:
    chunks: list[str] = []
    pages = 0
    for path in paths:
        if path.suffix == ".pdf":
            with _pdf(path, password) as doc:
                pages += len(doc)
                if pages > max_pages:
                    raise DocumentError("pages")
                for page in doc:
                    chunks.append(page.get_text())
                    if sum(map(len, chunks)) > MAX_TEXT_CHARS:
                        raise DocumentError("large")
        elif path.suffix == ".docx":
            try:
                # Refuse ZIP bombs before python-docx inflates XML and embedded media.
                with zipfile.ZipFile(path) as package:
                    infos = package.infolist()
                    if len(infos) > 2000 or sum(info.file_size for info in infos) > 50 * 1024 * 1024:
                        raise DocumentError("large")
                docx = Document(str(path))
                chunks.extend(p.text for p in docx.paragraphs)
                for table in docx.tables:
                    chunks.extend(cell.text for row in table.rows for cell in row.cells)
            except (ValueError, KeyError, zipfile.BadZipFile) as e:
                raise DocumentError() from e
        elif path.suffix == ".txt":
            chunks.append(path.read_text(encoding="utf-8", errors="replace"))
        else:
            pages += 1
            if pages > max_pages:
                raise DocumentError("pages")
        if sum(map(len, chunks)) > MAX_TEXT_CHARS:
            raise DocumentError("large")
    return text_counts("\n".join(chunks), pages)


def ocr_images(paths: list[Path], outdir: Path, *, max_pages: int, password: str = "") -> list[Path]:
    """Render bounded PDF pages or normalize images for either OCR backend."""
    images: list[Path] = []
    for path in paths:
        if path.suffix == ".pdf":
            with _pdf(path, password) as doc:
                if len(images) + len(doc) > max_pages:
                    raise DocumentError("pages")
                for page in doc:
                    target = outdir / f"ocr-{len(images) + 1:04}.png"
                    target.write_bytes(_render(page, 2))
                    images.append(target)
        elif path.suffix not in {".docx", ".txt"}:
            if len(images) >= max_pages:
                raise DocumentError("pages")
            target = outdir / f"ocr-{len(images) + 1:04}.png"
            image = _image(path)
            image.thumbnail((2400, 3400))
            image.save(target, "PNG")
            images.append(target)
    return images


def tesseract_text(images: list[Path], *, command: str = "tesseract", timeout: int = 45) -> str:
    executable = shutil.which(command)
    if not executable:
        raise DocumentError("unavailable")
    chunks: list[str] = []
    for image in images:
        try:
            result = subprocess.run(
                [executable, str(image), "stdout", "-l", "ara+eng", "--psm", "3"],
                capture_output=True, timeout=timeout, check=True,
            )
        except (subprocess.SubprocessError, OSError) as e:
            raise DocumentError("unavailable") from e
        chunks.append(result.stdout.decode("utf-8", errors="replace").strip())
        if sum(map(len, chunks)) > MAX_TEXT_CHARS:
            raise DocumentError("large")
    return "\n\n".join(chunks)


def export_text(text: str, output: Path) -> Path:
    if output.suffix == ".docx":
        document = Document()
        for paragraph in text.split("\n"):
            document.add_paragraph(paragraph)
        document.save(str(output))
    else:
        output.write_text(text, encoding="utf-8")
    return output


def _compress_images(doc: Any) -> None:
    seen: set[int] = set()
    for page in doc:
        for info in page.get_images(full=True):
            xref = info[0]
            if xref in seen or info[1]:  # avoid replacing images with a soft mask
                continue
            seen.add(xref)
            try:
                image_data = doc.extract_image(xref)
                with Image.open(io.BytesIO(image_data["image"])) as image:
                    if image.width * image.height > MAX_IMAGE_PIXELS:
                        continue
                    normalized = image.convert("RGB")
                    normalized.thumbnail((1600, 1600))
                    stream = io.BytesIO()
                    normalized.save(stream, format="JPEG", quality=72, optimize=True)
                    if len(stream.getvalue()) < len(image_data["image"]):
                        page.replace_image(xref, stream=stream.getvalue())
            except (RuntimeError, ValueError, OSError):
                continue


def process(
    operation: str, paths: list[Path], outdir: Path, *, options: dict[str, Any] | None = None,
    max_pages: int = 50,
) -> OperationResult:
    options = options or {}
    if not paths:
        raise DocumentError()
    if operation == "count":
        return extract_text(paths, max_pages=max_pages, password=str(options.get("password", "")))
    if operation == "image_pdf":
        if len(paths) > max_pages:
            raise DocumentError("pages")
        return OperationResult(files=[images_to_pdf(paths, outdir / "images.pdf")], pages=len(paths))
    if operation == "enhance":
        if len(paths) > max_pages:
            raise DocumentError("pages")
        return OperationResult(files=[enhance_scan(path, outdir / f"scan-{i:03}.png") for i, path in enumerate(paths, 1)], pages=len(paths))
    if any(path.suffix != ".pdf" for path in paths):
        raise DocumentError()
    total = 0
    for path in paths:
        with _pdf(path, str(options.get("password", ""))) as source:
            total += len(source)
            if total > max_pages:
                raise DocumentError("pages")
    if operation == "merge":
        with fitz.open() as result:
            for path in paths:
                with _pdf(path, str(options.get("password", ""))) as source:
                    result.insert_pdf(source)
            target = outdir / "merged.pdf"
            result.save(target, garbage=4, deflate=True)
        return OperationResult(files=[target], pages=total)
    if len(paths) != 1:
        raise DocumentError()
    with _pdf(paths[0], str(options.get("password", ""))) as doc:
        targets: list[Path] = []
        if operation in {"split", "extract"}:
            value = str(options.get("ranges", ""))
            groups = value.split(";") if operation == "split" else [value]
            if operation == "split" and not value:
                groups = [str(i) for i in range(1, len(doc) + 1)]
            selections = [parse_ranges(group, len(doc)) for group in groups]
            if sum(map(len, selections)) > max_pages:
                raise DocumentError("pages")
            for i, pages in enumerate(selections, 1):
                with fitz.open() as extracted:
                    for page in pages:
                        extracted.insert_pdf(doc, from_page=page, to_page=page)
                    target = outdir / f"{operation}-{i:03}.pdf"
                    extracted.save(target, garbage=4, deflate=True)
                    targets.append(target)
        elif operation == "images":
            for i, page in enumerate(doc, 1):
                data = _render(page)
                if options.get("format") == "album":
                    target = outdir / f"page-{i:04}.jpg"
                    with Image.open(io.BytesIO(data)) as rendered:
                        image = rendered.convert("RGB")
                        image.thumbnail((2000, 2000))
                        if image.width > image.height * 20 or image.height > image.width * 20:
                            image = ImageOps.pad(image, (max(image.width, image.height // 20 + 1),
                                                       max(image.height, image.width // 20 + 1)), color="white")
                        image.save(target, "JPEG", quality=85)
                else:
                    target = outdir / f"page-{i:04}.png"
                    target.write_bytes(data)
                targets.append(target)
            if options.get("format", "zip") != "album":
                archive = outdir / "pages.zip"
                with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zipped:
                    for target in targets:
                        zipped.write(target, target.name)
                targets = [archive]
        else:
            save_options: dict[str, Any] = {"garbage": 4, "deflate": True}
            if operation == "compress":
                _compress_images(doc)
            elif operation == "rotate":
                angle = int(options.get("angle", 90))
                if angle not in {90, 180, 270}:
                    raise DocumentError()
                for page in doc:
                    page.set_rotation((page.rotation + angle) % 360)
            elif operation == "numbers":
                for i, page in enumerate(doc, 1):
                    page.insert_textbox(fitz.Rect(20, page.rect.height - 28, page.rect.width - 20, page.rect.height - 10), str(i), fontsize=10, align=1)
            elif operation == "watermark":
                text = str(options.get("watermark", "")).strip()
                if not text or len(text) > 200:
                    raise DocumentError()
                for page in doc:
                    # Story uses built-in Unicode fonts, preserving Arabic text.
                    import html
                    rect = fitz.Rect(24, page.rect.height / 2 - 25, page.rect.width - 24, page.rect.height / 2 + 40)
                    page.insert_htmlbox(rect, f'<div style="text-align:center;color:#aaaaaa;font-size:24px">{html.escape(text)}</div>', opacity=0.25)
            elif operation == "encrypt":
                password = str(options.get("password", ""))
                if not 1 <= len(password.encode("utf-8")) <= 127:
                    raise DocumentError("password")
                # These runtime constants are absent from PyMuPDF's distributed type stubs.
                save_options.update(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password)  # type: ignore[attr-defined]
            elif operation == "decrypt":
                save_options["encryption"] = fitz.PDF_ENCRYPT_NONE  # type: ignore[attr-defined]
            else:
                raise DocumentError()
            target = outdir / f"{operation}.pdf"
            doc.save(target, **save_options)
            targets.append(target)
    return OperationResult(files=targets, pages=total)
